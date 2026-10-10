import asyncio
import os
import subprocess
import sys
from datetime import timedelta
from io import StringIO
from pathlib import Path
from unittest.mock import AsyncMock, Mock, patch

from config.client_ip import resolve_client_ip
from config.egress import playwright_launch_kwargs
from config.proxy_headers import TrustedProxyHeadersMiddleware
from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import RequestFactory, SimpleTestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document
from rest_framework import status
from rest_framework.test import APITestCase

from apps.billing.models import CoinTransaction, CoinWallet
from apps.scans.crawler import (
    _close_playwright_resources,
    _enforce_public_request,
    _enforce_public_websocket,
    _make_context,
    classify_blocked,
    classify_cf_challenge,
    compute_min_interval,
)
from apps.scans.models import AuthorizationConsent, Finding, Page, ScanJob, VerifiedDomain
from apps.scans.reports import get_severity_display, mask_pii_evidence, render_report_docx
from apps.scans.scanners import (
    PageAnalysisInput,
    analyze_aeo,
    analyze_data_exposure,
    analyze_geo_fast,
    analyze_page,
    analyze_security_site_level,
    analyze_site_signals,
    calculate_scores,
    detect_faq_structure,
    detect_pii_in_text,
    is_admin_path,
    is_binary_resource,
    is_valid_luhn,
    is_valid_tw_national_id,
    make_finding,
    parse_html_signals,
)
from apps.scans.serializers import FindingSerializer
from apps.scans.services import PublicScanTargetError, assert_public_http_url


class ScanJobModelTests(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="tester",
            email="tester@example.com",
            password="safe-test-password",
        )

    def test_active_scan_requires_extra_authorization(self):
        scan_job = ScanJob(
            user=self.user,
            original_url="https://example.com/",
            normalized_url="https://example.com/",
            origin="https://example.com",
            scan_mode=ScanJob.ScanMode.ACTIVE,
            active_testing_authorized=False,
        )

        with self.assertRaisesMessage(Exception, "主動測試必須先取得額外授權。"):
            scan_job.clean()

    def test_passive_scan_defaults_are_safe(self):
        scan_job = ScanJob.objects.create(
            user=self.user,
            original_url="https://example.com/",
            normalized_url="https://example.com/",
            origin="https://example.com",
        )

        self.assertEqual(scan_job.scan_mode, ScanJob.ScanMode.PASSIVE)
        self.assertEqual(scan_job.max_depth, 3)
        self.assertEqual(scan_job.max_pages, 50)
        self.assertTrue(scan_job.respect_robots)


class StaticScannerTests(APITestCase):
    def test_analyze_page_returns_seo_geo_and_security_findings(self):
        page_input = PageAnalysisInput(
            url="http://example.com/",
            final_url="http://example.com/",
            title="短",
            html=(
                "<html><body><h1>主標題</h1><img src='/a.png'>"
                "<form method='post'></form></body></html>"
            ),
            headers={},
            element_boxes={"h1": {"x": 1, "y": 2, "width": 3, "height": 4}},
        )

        findings = analyze_page(page_input)
        categories = {finding["category"] for finding in findings}
        titles = {finding["title"] for finding in findings}

        self.assertIn("seo", categories)
        self.assertIn("geo", categories)
        self.assertIn("security", categories)
        self.assertIn("Meta title 長度不理想", titles)
        self.assertIn("可補充 JSON-LD 結構化資料", titles)
        # HTTPS/HSTS/CSP 等站台層級標頭檢查已搬到 analyze_security_site_level()
        # （見下方測試），analyze_page() 只保留逐頁各自判斷的 CSRF 檢查。
        self.assertIn("表單可能缺少 CSRF token", titles)

    def test_analyze_security_site_level_flags_missing_https_and_headers(self):
        pages = [
            {
                "final_url": "http://example.com/",
                "url": "http://example.com/",
                "headers": {"content-type": "text/html"},
            }
        ]

        findings = analyze_security_site_level(pages)
        titles = {finding["title"] for finding in findings}

        self.assertIn("頁面未使用 HTTPS", titles)
        self.assertIn("缺少 HSTS", titles)
        self.assertIn("缺少 CSP", titles)

    def test_analyze_security_site_level_only_evaluates_first_page_with_headers(self):
        # 第一頁 headers 為空（爬蟲可能因逾時等原因沒抓到），應跳過改看下一頁；
        # 找到的頁面若已具備全部必要標頭，就不該產生任何 finding——
        # 驗證「整批頁面只評估一次」而非逐頁重複判斷。
        pages = [
            {"final_url": "https://example.com/a", "url": "https://example.com/a", "headers": {}},
            {
                "final_url": "https://example.com/b",
                "url": "https://example.com/b",
                "headers": {
                    "strict-transport-security": "max-age=1",
                    "content-security-policy": "default-src 'self'",
                    "x-frame-options": "DENY",
                    "x-content-type-options": "nosniff",
                },
            },
        ]

        findings = analyze_security_site_level(pages)

        self.assertEqual(findings, [])

    def test_analyze_security_site_level_skips_blocked_pages(self):
        # 第一頁被標記 blocked_reason（例如跨網域導向、CF challenge）：它的 headers
        # 來自不受信任的來源，不能代表使用者自己網站的設定，必須跳過改看下一頁。
        pages = [
            {
                "final_url": "https://third-party.example/",
                "url": "https://example.com/",
                "headers": {"content-type": "text/html"},
                "blocked_reason": "跨網域導向，超出授權範圍",
            },
            {
                "final_url": "https://example.com/b",
                "url": "https://example.com/b",
                "headers": {
                    "strict-transport-security": "max-age=1",
                    "content-security-policy": "default-src 'self'",
                    "x-frame-options": "DENY",
                    "x-content-type-options": "nosniff",
                },
                "blocked_reason": "",
            },
        ]

        findings = analyze_security_site_level(pages)

        self.assertEqual(findings, [])

    def test_calculate_scores_returns_top_actions_without_code(self):
        findings = [
            {
                "category": "security",
                "severity": "high",
                "title": "頁面未使用 HTTPS",
                "priority_score": 90,
            },
            {
                "category": "seo",
                "severity": "low",
                "title": "缺少 canonical URL",
                "priority_score": 30,
            },
        ]

        overall_score, category_scores, top_actions = calculate_scores(findings)

        self.assertLess(overall_score, 100)
        self.assertLess(category_scores["security"], category_scores["aeo"])
        self.assertEqual(top_actions[0]["title"], "頁面未使用 HTTPS")

    def test_make_finding_adds_evidence_first_metadata(self):
        finding = make_finding(
            category=Finding.Category.SEO,
            severity=Finding.Severity.MEDIUM,
            title="缺少 meta description",
            description="頁面未提供摘要。",
            remediation="補上可描述頁面內容的 meta description。",
            evidence="未找到 meta[name='description']",
            evidence_type="html_rule",
            evidence_source="seo_rule_engine",
        )

        self.assertTrue(finding["rule_id"].startswith("SEO_"))
        self.assertEqual(finding["evidence_type"], "html_rule")
        self.assertEqual(finding["evidence_source"], "seo_rule_engine")
        self.assertEqual(finding["evidence_json"]["excerpt"], "未找到 meta[name='description']")
        self.assertEqual(finding["ai_explanation"], "")
        self.assertEqual(finding["ai_remediation"], "")

    def test_finding_serializer_exposes_evidence_first_fields(self):
        user = get_user_model().objects.create_user(
            username="evidence-user",
            email="evidence@example.com",
            password="safe-test-password",
        )
        scan_job = ScanJob.objects.create(
            user=user,
            original_url="https://example.com/",
            normalized_url="https://example.com/",
            origin="https://example.com",
        )
        finding_payload = make_finding(
            category=Finding.Category.SECURITY,
            severity=Finding.Severity.HIGH,
            title="頁面未使用 HTTPS",
            description="目標頁面使用 HTTP。",
            remediation="改用 HTTPS 並設定 HSTS。",
            evidence="scheme=http",
            rule_id="SECURITY_HTTPS_REQUIRED",
            evidence_type="url_scheme",
            evidence_source="security_rule_engine",
        )
        finding = Finding.objects.create(scan_job=scan_job, **finding_payload)

        data = FindingSerializer(finding).data

        self.assertEqual(data["rule_id"], "SECURITY_HTTPS_REQUIRED")
        self.assertEqual(data["evidence_type"], "url_scheme")
        self.assertEqual(data["evidence_source"], "security_rule_engine")
        self.assertEqual(data["evidence_json"]["excerpt"], "scheme=http")

    def test_report_handles_unknown_action_severity(self):
        user = get_user_model().objects.create_user(
            username="report-user",
            email="report@example.com",
            password="safe-test-password",
        )
        scan_job = ScanJob.objects.create(
            user=user,
            original_url="https://example.com/",
            normalized_url="https://example.com/",
            origin="https://example.com",
            status=ScanJob.Status.COMPLETED,
            top_actions=[
                {
                    "category": "security",
                    "severity": "warning",
                    "title": "未知嚴重度測試",
                }
            ],
        )

        self.assertEqual(get_severity_display("warning"), "warning")
        self.assertTrue(render_report_docx(scan_job).endswith(".docx"))

    def test_report_masks_pii_evidence_even_when_truncated(self):
        # 60 筆信用卡號串接，總長度超過報告的 1000 字截斷點，最後一筆極可能被
        # 切在數字中間；遮罩必須在完整字串上跑完才截斷，順序顛倒的話殘缺數字
        # 長度不足以命中 regex，會以明文殘留在 .docx 裡。
        user = get_user_model().objects.create_user(
            username="pii-report-user",
            email="pii-report@example.com",
            password="safe-test-password",
        )
        scan_job = ScanJob.objects.create(
            user=user,
            original_url="https://example.com/",
            normalized_url="https://example.com/",
            origin="https://example.com",
            status=ScanJob.Status.COMPLETED,
        )
        card_numbers = [f"41111111{i:08d}" for i in range(60)]
        evidence = "信用卡號（60 筆）：" + ", ".join(card_numbers)
        finding_payload = make_finding(
            category=Finding.Category.SECURITY,
            severity=Finding.Severity.HIGH,
            title="頁面外洩個人資料 (PII)",
            description="測試用",
            remediation="測試用",
            evidence=evidence,
        )
        Finding.objects.create(scan_job=scan_job, **finding_payload)

        output_path = render_report_docx(scan_job)
        document = Document(output_path)
        full_text = "\n".join(p.text for p in document.paragraphs)

        for card in card_numbers:
            self.assertNotIn(card, full_text)

    def test_report_masks_pii_from_exposure_scanner_findings_too(self):
        # security/exposure_scanner.py 產生的敏感檔案外洩 finding 用不同的 rule_id
        # （不是 SECURITY_PII_* 前綴），但 evidence 的「檔案內容片段」一樣可能含
        # 原始個資；遮罩不能只靠 rule_id 白名單判斷是不是 PII finding。
        user = get_user_model().objects.create_user(
            username="exposure-report-user",
            email="exposure-report@example.com",
            password="safe-test-password",
        )
        scan_job = ScanJob.objects.create(
            user=user,
            original_url="https://example.com/",
            normalized_url="https://example.com/",
            origin="https://example.com",
            status=ScanJob.Status.COMPLETED,
        )
        finding_payload = make_finding(
            category=Finding.Category.SECURITY,
            severity=Finding.Severity.HIGH,
            rule_id="exposure-env-file",
            title="敏感檔案外洩：環境變數檔",
            description="測試用",
            remediation="測試用",
            evidence="檔案內容片段：\nADMIN_EMAIL=admin@example.com\nADMIN_PHONE=0912345678",
        )
        Finding.objects.create(scan_job=scan_job, **finding_payload)

        output_path = render_report_docx(scan_job)
        document = Document(output_path)
        full_text = "\n".join(p.text for p in document.paragraphs)

        self.assertNotIn("admin@example.com", full_text)
        self.assertNotIn("0912345678", full_text)

    def test_mask_pii_evidence_keeps_edges_only(self):
        masked = mask_pii_evidence("聯絡信箱：service@example.com，電話：0987654321")

        self.assertIn("se*****@example.com", masked)
        self.assertIn("09******21", masked)
        self.assertNotIn("service@example.com", masked)
        self.assertNotIn("0987654321", masked)


@override_settings(ARGUS_AUTO_QUEUE_SCANS=False)
class ScanJobApiTests(APITestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="api-user",
            email="api@example.com",
            password="safe-test-password",
        )
        # 預先把測試使用者的 coin 加滿，避免 coin 不足干擾建立掃描的行為測試
        CoinWallet.objects.filter(user=self.user).update(balance=10000)
        self.client.force_authenticate(self.user)
        self.url = reverse("scan-list")

    def test_create_scan_requires_authorization_confirmation(self):
        response = self.client.post(
            self.url,
            {
                "url": "https://example.com/",
                "authorization_confirmed": False,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(ScanJob.objects.count(), 0)

    def test_create_scan_records_authorization_consent(self):
        response = self.client.post(
            self.url,
            {
                "url": "example.com",
                "authorization_confirmed": True,
            },
            format="json",
            HTTP_USER_AGENT="ArgusTest/1.0",
            REMOTE_ADDR="203.0.113.10",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        scan_job = ScanJob.objects.get()
        consent = AuthorizationConsent.objects.get(scan_job=scan_job)
        self.assertEqual(scan_job.normalized_url, "https://example.com/")
        self.assertEqual(scan_job.origin, "https://example.com")
        self.assertEqual(consent.authorized_domain, "example.com")
        self.assertEqual(consent.ip_address, "203.0.113.10")
        self.assertEqual(consent.user_agent, "ArgusTest/1.0")

    def test_create_scan_response_uses_read_model_shape(self):
        response = self.client.post(
            self.url,
            {
                "url": "https://example.com/",
                "authorization_confirmed": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertIn("status", response.data)
        self.assertIn("normalized_url", response.data)
        self.assertNotIn("authorization_confirmed", response.data)

    def test_create_scan_without_categories_defaults_to_all(self):
        response = self.client.post(
            self.url,
            {"url": "https://example.com/", "authorization_confirmed": True},
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(
            set(ScanJob.objects.get().effective_categories),
            {"seo", "aeo", "geo", "ux", "security"},
        )

    def test_create_scan_partial_categories_stored_and_charged(self):
        response = self.client.post(
            self.url,
            {
                "url": "https://example.com/",
                "authorization_confirmed": True,
                "categories": ["seo", "geo"],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        scan_job = ScanJob.objects.get()
        self.assertEqual(scan_job.categories, ["seo", "geo"])
        # 預設 max_pages=50：50 頁 × 2 維 × 2 coin = 200
        hold = scan_job.coin_transactions.get(kind=CoinTransaction.Kind.SCAN_HOLD)
        self.assertEqual(hold.amount, -200)

    def test_create_scan_rejects_empty_categories(self):
        response = self.client.post(
            self.url,
            {
                "url": "https://example.com/",
                "authorization_confirmed": True,
                "categories": [],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(ScanJob.objects.count(), 0)

    def test_active_mode_requires_security_category(self):
        # 未通過網域驗證會先撞網域閘門，故先建已驗證網域讓測試聚焦在維度規則
        VerifiedDomain.objects.create(
            user=self.user,
            domain="example.com",
            token="0f1e2d3c4b5a69788796a5b4c3d2e1f0",
            status=VerifiedDomain.Status.VERIFIED,
            method=VerifiedDomain.Method.DNS_TXT,
            verified_at=timezone.now(),
            expires_at=timezone.now() + timedelta(days=90),
        )
        response = self.client.post(
            self.url,
            {
                "url": "https://example.com/",
                "authorization_confirmed": True,
                "scan_mode": "active",
                "active_testing_authorized": True,
                "categories": ["seo", "geo"],
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("資安", str(response.data))
        self.assertEqual(ScanJob.objects.count(), 0)

    def test_model_clean_rejects_active_without_security(self):
        scan_job = ScanJob(
            user=self.user,
            original_url="https://example.com/",
            normalized_url="https://example.com/",
            origin="https://example.com",
            scan_mode=ScanJob.ScanMode.ACTIVE,
            active_testing_authorized=True,
            categories=["seo"],
        )
        with self.assertRaises(ValidationError):
            scan_job.clean()

    def test_obvious_third_party_requires_reconfirmation(self):
        response = self.client.post(
            self.url,
            {
                "url": "https://google.com/",
                "authorization_confirmed": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("third_party_reconfirmed", response.data)

    def test_active_scan_requires_active_authorization(self):
        response = self.client.post(
            self.url,
            {
                "url": "https://example.com/",
                "authorization_confirmed": True,
                "scan_mode": ScanJob.ScanMode.ACTIVE,
                "active_testing_authorized": False,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(ScanJob.objects.count(), 0)

    def test_active_scan_with_extra_authorization_is_recorded(self):
        # 主動測試閘門：目標網域必須先通過所有權驗證（2026-09 新增的技術性驗證）
        from datetime import timedelta

        from django.utils import timezone

        VerifiedDomain.objects.create(
            user=self.user,
            domain="example.com",
            token="a" * 32,
            status=VerifiedDomain.Status.VERIFIED,
            method=VerifiedDomain.Method.DNS_TXT,
            verified_at=timezone.now(),
            expires_at=timezone.now() + timedelta(days=90),
        )
        response = self.client.post(
            self.url,
            {
                "url": "https://example.com/",
                "authorization_confirmed": True,
                "scan_mode": ScanJob.ScanMode.ACTIVE,
                "active_testing_authorized": True,
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        scan_job = ScanJob.objects.get()
        consent = AuthorizationConsent.objects.get(scan_job=scan_job)
        self.assertEqual(scan_job.scan_mode, ScanJob.ScanMode.ACTIVE)
        self.assertTrue(scan_job.active_testing_authorized)
        self.assertTrue(consent.active_testing_authorized)

    def test_status_endpoint_returns_scan_status(self):
        scan_job = ScanJob.objects.create(
            user=self.user,
            original_url="https://example.com/",
            normalized_url="https://example.com/",
            origin="https://example.com",
        )

        response = self.client.get(reverse("scan-status", args=[scan_job.id]))

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["status"], ScanJob.Status.QUEUED)


class CrawlerHelperTests(APITestCase):
    def test_playwright_cleanup_error_is_best_effort(self):
        broken_context = Mock()
        broken_context.close = AsyncMock(
            side_effect=RuntimeError("test-only cleanup failure")
        )
        browser = Mock()
        browser.close = AsyncMock()

        asyncio.run(_close_playwright_resources(broken_context, browser))

        broken_context.close.assert_awaited_once()
        browser.close.assert_awaited_once()

    def test_compute_min_interval_active_enforces_rps_cap(self):
        # 主動模式 RPS=2，兩次請求至少間隔 0.5 秒
        interval = compute_min_interval("active", active_rps=2, passive_rps=5)
        self.assertEqual(interval, 0.5)

    def test_compute_min_interval_passive_is_faster_than_active(self):
        active = compute_min_interval("active", active_rps=2, passive_rps=5)
        passive = compute_min_interval("passive", active_rps=2, passive_rps=5)
        self.assertLess(passive, active)

    def test_classify_blocked_flags_forbidden_and_rate_limited(self):
        self.assertNotEqual(classify_blocked(401), "")
        self.assertNotEqual(classify_blocked(403), "")
        self.assertNotEqual(classify_blocked(429), "")

    def test_classify_blocked_allows_normal_status(self):
        self.assertEqual(classify_blocked(200), "")
        self.assertEqual(classify_blocked(404), "")
        self.assertEqual(classify_blocked(None), "")

    def test_classify_cf_challenge_detects_js_challenge(self):
        html = (
            "<html><head><title>Just a moment...</title>"
            '<script src="/cdn-cgi/challenge-platform/h/b/orchestrate/chl_page/v1"></script>'
            "</head><body></body></html>"
        )
        self.assertEqual(
            classify_cf_challenge(html),
            "Cloudflare JavaScript 驗證，自動掃描無法通過",
        )

    def test_classify_cf_challenge_detects_browser_verification(self):
        html = '<div id="cf-browser-verification">驗證中</div>'
        self.assertEqual(
            classify_cf_challenge(html),
            "Cloudflare JavaScript 驗證，自動掃描無法通過",
        )

    def test_classify_cf_challenge_detects_turnstile(self):
        html = '<div class="cf-turnstile" data-sitekey="xxx"></div>'
        self.assertEqual(
            classify_cf_challenge(html),
            "Cloudflare Turnstile 驗證，需使用者互動才能通過",
        )

    def test_classify_cf_challenge_turnstile_takes_priority_over_js(self):
        # 兩種標記並存時 Turnstile 較嚴重（必擋），應優先回報
        html = (
            '<script src="/cdn-cgi/challenge-platform/x.js"></script>'
            '<div class="cf-turnstile"></div>'
        )
        self.assertEqual(
            classify_cf_challenge(html),
            "Cloudflare Turnstile 驗證，需使用者互動才能通過",
        )

    def test_classify_cf_challenge_ignores_normal_html(self):
        html = "<html><body><h1>歡迎</h1><p>這是正常頁面內容</p></body></html>"
        self.assertEqual(classify_cf_challenge(html), "")

    def test_classify_cf_challenge_handles_empty_input(self):
        self.assertEqual(classify_cf_challenge(""), "")
        self.assertEqual(classify_cf_challenge(None), "")

    def test_classify_cf_challenge_ignores_bot_detection_script_on_normal_page(self):
        # CF Bot 偵測在正常頁面尾端插入的背景腳本（實際取自 ntubimdbirc.tw 首頁）不是攔截頁
        html = (
            "<html><head><title>NTUB BIRC</title></head><body><a href=\"/about\">關於</a>"
            "<script>window.__CF$cv$params={r:'a4454a895a20420b'};"
            "(function(){var s=document.createElement('script');"
            "s.src='/cdn-cgi/challenge-platform/scripts/precursor/main.js';"
            "document.head.appendChild(s);})();</script>"
            '<script src="/cdn-cgi/challenge-platform/scripts/jsd/main.js"></script></body></html>'
        )
        self.assertEqual(classify_cf_challenge(html), "")

    def test_classify_cf_challenge_detects_challenge_options(self):
        html = "<script>window._cf_chl_opt={cvId: '3'};</script>"
        self.assertEqual(
            classify_cf_challenge(html),
            "Cloudflare JavaScript 驗證，自動掃描無法通過",
        )

    def test_classify_cf_challenge_ignores_just_a_moment_phrase(self):
        # 不收錄「Just a moment」短語標記，避免正常文章正文出現該短語時誤判
        html = "<p>Just a moment, please wait while I finish typing...</p>"
        self.assertEqual(classify_cf_challenge(html), "")


class ScanTargetPolicyTests(APITestCase):
    @patch(
        "apps.scans.services.socket.getaddrinfo",
        return_value=[(None, None, None, "", ("93.184.216.34", 0))],
    )
    def test_public_domain_is_normalized(self, _mock_getaddrinfo):
        self.assertEqual(
            assert_public_http_url("HTTPS://Example.COM/path?q=1"),
            "https://example.com/path?q=1",
        )

    def test_private_and_metadata_addresses_are_rejected(self):
        for target in (
            "http://127.0.0.1/",
            "http://[::1]/",
            "http://10.0.0.1/",
            "http://172.16.0.1/",
            "http://192.168.0.1/",
            "http://169.254.169.254/latest/meta-data/",
        ):
            with self.subTest(target=target):
                with self.assertRaises(PublicScanTargetError):
                    assert_public_http_url(target)

    @patch(
        "apps.scans.services.socket.getaddrinfo",
        return_value=[
            (None, None, None, "", ("93.184.216.34", 0)),
            (None, None, None, "", ("10.0.0.8", 0)),
        ],
    )
    def test_domain_with_any_private_dns_answer_is_rejected(self, _mock_getaddrinfo):
        with self.assertRaises(PublicScanTargetError):
            assert_public_http_url("https://mixed.example/")

    def test_userinfo_and_non_web_ports_are_rejected(self):
        for target in (
            "https://user:password@example.com/",
            "https://example.com:22/",
        ):
            with self.subTest(target=target):
                with self.assertRaises(PublicScanTargetError):
                    assert_public_http_url(target)

    async def test_playwright_route_aborts_private_request_before_network(self):
        route = AsyncMock()
        request = Mock(url="http://127.0.0.1/private")

        await _enforce_public_request(route, request)

        route.abort.assert_awaited_once_with("blockedbyclient")
        route.continue_.assert_not_awaited()

    @patch(
        "apps.scans.services.socket.getaddrinfo",
        return_value=[(None, None, None, "", ("93.184.216.34", 0))],
    )
    async def test_playwright_route_allows_public_request(self, _mock_getaddrinfo):
        route = AsyncMock()
        request = Mock(url="https://cdn.example.com/app.js")

        await _enforce_public_request(route, request)

        route.continue_.assert_awaited_once_with()
        route.abort.assert_not_awaited()

    @patch(
        "apps.scans.services.socket.getaddrinfo",
        return_value=[(None, None, None, "", ("93.184.216.34", 0))],
    )
    async def test_playwright_route_blocks_cross_origin_main_navigation(
        self,
        _mock_getaddrinfo,
    ):
        route = AsyncMock()
        page = Mock()
        frame = Mock()
        page.main_frame = frame
        frame.page = page
        request = Mock(url="https://other.example/redirected", frame=frame)
        request.is_navigation_request.return_value = True

        await _enforce_public_request(route, request, "https://example.com")

        route.abort.assert_awaited_once_with("blockedbyclient")
        route.continue_.assert_not_awaited()

    async def test_playwright_websocket_closes_private_target(self):
        websocket_route = Mock(url="ws://169.254.169.254/socket")
        websocket_route.close = AsyncMock()

        await _enforce_public_websocket(websocket_route)

        websocket_route.close.assert_awaited_once()
        websocket_route.connect_to_server.assert_not_called()

    @patch(
        "apps.scans.services.socket.getaddrinfo",
        return_value=[(None, None, None, "", ("93.184.216.34", 0))],
    )
    async def test_playwright_websocket_closes_cross_origin_target(
        self,
        _mock_getaddrinfo,
    ):
        websocket_route = Mock(url="wss://other.example/socket")
        websocket_route.close = AsyncMock()

        await _enforce_public_websocket(websocket_route, "https://example.com")

        websocket_route.close.assert_awaited_once()
        websocket_route.connect_to_server.assert_not_called()


class EgressSettingsTests(SimpleTestCase):
    def _settings_process(self, proxy_url: str):
        env = os.environ.copy()
        env.update(
            {
                "DJANGO_SECRET_KEY": "test-only-django-secret-with-at-least-32-bytes",
                "PASSWORD_RESET_TOKEN_PEPPER": "test-only-reset-pepper-with-at-least-32-bytes",
                "ARGUS_PAYMENT_MODE": "disabled",
                "ARGUS_EGRESS_PROXY_URL": proxy_url,
                "NO_PROXY": "*",
                "no_proxy": "*",
            }
        )
        return subprocess.run(
            [
                sys.executable,
                "-c",
                (
                    "import os; from config import settings; "
                    "assert os.environ['HTTP_PROXY'] == settings.ARGUS_EGRESS_PROXY_URL; "
                    "assert os.environ['http_proxy'] == settings.ARGUS_EGRESS_PROXY_URL; "
                    "assert os.environ['NO_PROXY'] != '*'; "
                    "assert os.environ['no_proxy'] == os.environ['NO_PROXY']"
                ),
            ],
            cwd=Path(__file__).resolve().parents[2],
            env=env,
            capture_output=True,
            text=True,
            timeout=20,
            check=False,
        )

    def test_proxy_overrides_upper_and_lowercase_environment(self):
        result = self._settings_process("http://egress-proxy:3128")
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_proxy_rejects_even_empty_userinfo(self):
        result = self._settings_process("http://:@egress-proxy:3128")
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("ARGUS_EGRESS_PROXY_URL", result.stderr)


class TrustedProxyClientIPTests(APITestCase):
    @override_settings(TRUSTED_PROXY_CIDRS=[])
    def test_direct_client_cannot_spoof_forwarded_header(self):
        request = Mock(META={
            "REMOTE_ADDR": "203.0.113.10",
            "HTTP_X_FORWARDED_FOR": "1.2.3.4",
        })
        self.assertEqual(resolve_client_ip(request), "203.0.113.10")

    @override_settings(TRUSTED_PROXY_CIDRS=["10.0.0.0/8", "2001:db8:1::/48"])
    def test_trusted_proxy_chain_uses_first_untrusted_hop_from_right(self):
        request = Mock(META={
            "REMOTE_ADDR": "10.0.0.5",
            "HTTP_X_FORWARDED_FOR": "198.51.100.8, 203.0.113.9, 10.0.0.4",
        })
        self.assertEqual(resolve_client_ip(request), "203.0.113.9")

    @override_settings(TRUSTED_PROXY_CIDRS=["10.0.0.0/8"])
    def test_malformed_or_excessive_chain_falls_back_to_remote(self):
        for forwarded in ("not-an-ip", ",".join(["1.1.1.1"] * 21)):
            with self.subTest(forwarded=forwarded):
                request = Mock(META={
                    "REMOTE_ADDR": "10.0.0.5",
                    "HTTP_X_FORWARDED_FOR": forwarded,
                })
                self.assertEqual(resolve_client_ip(request), "10.0.0.5")

    @override_settings(TRUSTED_PROXY_CIDRS=["2001:db8:1::/48"])
    def test_ipv6_forwarded_client_is_supported(self):
        request = Mock(META={
            "REMOTE_ADDR": "2001:db8:1::5",
            "HTTP_X_FORWARDED_FOR": "2001:4860:4860::8888",
        })
        self.assertEqual(resolve_client_ip(request), "2001:4860:4860::8888")

    @override_settings(
        TRUST_PROXY_SSL_HEADER=True,
        TRUSTED_PROXY_CIDRS=["10.0.0.0/8"],
        SECURE_PROXY_SSL_HEADER=("HTTP_X_FORWARDED_PROTO", "https"),
    )
    def test_direct_client_cannot_spoof_forwarded_proto(self):
        request = RequestFactory().get(
            "/",
            HTTP_X_FORWARDED_PROTO="https",
            REMOTE_ADDR="198.51.100.20",
        )
        middleware = TrustedProxyHeadersMiddleware(lambda req: req.is_secure())

        self.assertFalse(middleware(request))

    @override_settings(
        TRUST_PROXY_SSL_HEADER=True,
        TRUSTED_PROXY_CIDRS=["10.0.0.0/8"],
        SECURE_PROXY_SSL_HEADER=("HTTP_X_FORWARDED_PROTO", "https"),
    )
    def test_trusted_proxy_can_supply_forwarded_proto(self):
        request = RequestFactory().get(
            "/",
            HTTP_X_FORWARDED_PROTO="https",
            REMOTE_ADDR="10.0.0.5",
        )
        middleware = TrustedProxyHeadersMiddleware(lambda req: req.is_secure())

        self.assertTrue(middleware(request))

    @override_settings(ARGUS_EGRESS_PROXY_URL="")
    def test_playwright_proxy_is_omitted_by_default(self):
        self.assertEqual(playwright_launch_kwargs(), {})

    @override_settings(ARGUS_EGRESS_PROXY_URL="http://egress-proxy:3128")
    def test_playwright_proxy_is_applied_when_configured(self):
        self.assertEqual(
            playwright_launch_kwargs(),
            {"proxy": {"server": "http://egress-proxy:3128"}},
        )

    def test_crawler_context_blocks_service_workers_before_routing(self):
        context = Mock()
        context.route = AsyncMock()
        context.route_web_socket = AsyncMock()
        browser = Mock()
        browser.new_context = AsyncMock(return_value=context)

        created = asyncio.run(_make_context(browser, "https://example.com"))

        self.assertIs(created, context)
        browser.new_context.assert_awaited_once()
        self.assertEqual(
            browser.new_context.await_args.kwargs["service_workers"],
            "block",
        )
        context.route.assert_awaited_once()
        context.route_web_socket.assert_awaited_once()


class HealthEndpointTests(APITestCase):
    def test_liveness_and_readiness_are_public_and_healthy(self):
        live = self.client.get(reverse("health-live"))
        ready = self.client.get(reverse("health-ready"))

        self.assertEqual(live.status_code, status.HTTP_200_OK)
        self.assertEqual(ready.status_code, status.HTTP_200_OK)

    def test_favicon_is_served_as_static_asset_not_spa_html(self):
        response = self.client.get("/favicon.svg")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn("svg", response["Content-Type"])


class PiiDetectionTests(APITestCase):
    """PII 偵測測試：含檢查碼驗證、false positive 緩解、警示文字。"""

    # --- 身分證檢查碼 ---
    def test_is_valid_tw_national_id_accepts_known_valid(self):
        # 範例為自構符合內政部演算法的合法格式（非真實人物）
        self.assertTrue(is_valid_tw_national_id("A123456789"))
        self.assertTrue(is_valid_tw_national_id("B100000002"))

    def test_is_valid_tw_national_id_rejects_invalid_checksum(self):
        # 格式對但檢查碼錯誤
        self.assertFalse(is_valid_tw_national_id("A123456788"))
        self.assertFalse(is_valid_tw_national_id("A000000000"))

    def test_is_valid_tw_national_id_rejects_malformed(self):
        self.assertFalse(is_valid_tw_national_id(""))
        self.assertFalse(is_valid_tw_national_id("A12345678"))  # 太短
        self.assertFalse(is_valid_tw_national_id("A1234567890"))  # 太長
        self.assertFalse(is_valid_tw_national_id("1234567890"))  # 缺字母

    # --- Luhn 信用卡 ---
    def test_is_valid_luhn_accepts_known_test_cards(self):
        # Visa / MasterCard 測試卡號（業界公開測試用）
        self.assertTrue(is_valid_luhn("4111111111111111"))
        self.assertTrue(is_valid_luhn("5555555555554444"))

    def test_is_valid_luhn_rejects_invalid(self):
        self.assertFalse(is_valid_luhn("4111111111111112"))  # 末位錯
        self.assertFalse(is_valid_luhn("1234567890123456"))  # 隨機 16 位
        self.assertFalse(is_valid_luhn("123"))  # 太短

    # --- detect_pii_in_text 整合 ---
    def test_detect_pii_finds_email(self):
        result = detect_pii_in_text("聯絡 a@b.com 或 admin@example.tw 謝謝")
        self.assertIn("a@b.com", result["email"])
        self.assertIn("admin@example.tw", result["email"])

    def test_detect_pii_finds_taiwan_mobile(self):
        result = detect_pii_in_text("手機 0912345678 或 0987-654-321")
        self.assertIn("0912345678", result["mobile"])
        self.assertIn("0987-654-321", result["mobile"])

    def test_detect_pii_ignores_mobile_embedded_in_long_digits(self):
        # 09 開頭但是更長數字串的一部分（例如商品編號），不應誤判
        result = detect_pii_in_text("訂單號 12309123456789012")
        self.assertEqual(result["mobile"], [])

    def test_detect_pii_finds_national_id_with_valid_checksum(self):
        result = detect_pii_in_text("身分證 A123456789 已驗證")
        self.assertIn("A123456789", result["national_id"])

    def test_detect_pii_ignores_national_id_with_invalid_checksum(self):
        # 格式符合 regex 但檢查碼錯誤，應被過濾
        result = detect_pii_in_text("亂寫一個 A123456780 應該不通過")
        self.assertEqual(result["national_id"], [])

    def test_detect_pii_finds_credit_card_with_valid_luhn(self):
        result = detect_pii_in_text("卡號 4111-1111-1111-1111 已收")
        self.assertEqual(len(result["credit_card"]), 1)

    def test_detect_pii_ignores_random_digits_failing_luhn(self):
        result = detect_pii_in_text("流水號 1234567890123456 不是卡號")
        self.assertEqual(result["credit_card"], [])

    def test_detect_pii_ignores_bare_luhn_valid_without_context(self):
        # 通過 Luhn 但無分隔、附近也無信用卡關鍵字（流水號）→ 不採計，收斂誤報
        result = detect_pii_in_text("序號 4111111111111111 結束")
        self.assertEqual(result["credit_card"], [])

    def test_detect_pii_keeps_formatted_card_without_context(self):
        # 4-4-4-4 格式化卡號即使無關鍵字，仍視為高信心
        result = detect_pii_in_text("備註 4111-1111-1111-1111 完")
        self.assertEqual(result["credit_card"], ["4111-1111-1111-1111"])

    def test_detect_pii_keeps_bare_card_with_context(self):
        # 裸號但附近有「卡」關鍵字 → 採計
        result = detect_pii_in_text("信用卡號 4111111111111111")
        self.assertEqual(result["credit_card"], ["4111111111111111"])

    def test_card_needs_card_grouping_and_issuer(self):
        # 2026-10-10 實測誤報：日期＋流水號（8＋7 位）巧合通過 Luhn，不是卡片的分組方式
        self.assertEqual(detect_pii_in_text("檔案 20221027 0069652 結束")["credit_card"], [])
        # 位數與分組都對，但開頭不是任何卡組織（IIN）
        self.assertEqual(detect_pii_in_text("信用卡 1000-0000-0000-0008")["credit_card"], [])
        # Amex 4-6-5
        self.assertEqual(
            detect_pii_in_text("卡號 3782-822463-10005")["credit_card"], ["3782-822463-10005"]
        )

    def test_card_and_national_id_ignore_attribute_values_and_scripts(self):
        html_text = (
            '<html><body><img src="/_next/static/media/20221027 0069652-ISO 9001.jpg" '
            'alt="ISO">'
            '<a href="/files/A123456789.pdf">下載</a>'
            '<script>var cc = "4111-1111-1111-1111";</script>'
            "<p>認證證書</p></body></html>"
        )
        findings = analyze_data_exposure(self._page_input(html_text))
        self.assertFalse(any(f["rule_id"] == "SECURITY_PII_8B24BB8B28" for f in findings))
        # 同一個號碼寫在看得到的文字裡就會被抓到
        visible = analyze_data_exposure(
            self._page_input("<html><body><p>卡號 4111-1111-1111-1111</p></body></html>")
        )
        self.assertTrue(any(f["rule_id"] == "SECURITY_PII_8B24BB8B28" for f in visible))

    def test_detect_pii_dedups_repeated_values(self):
        result = detect_pii_in_text("a@b.com a@b.com a@b.com 重複出現")
        self.assertEqual(result["email"], ["a@b.com"])

    def test_detect_pii_empty_input(self):
        result = detect_pii_in_text("")
        self.assertEqual(result["email"], [])
        self.assertEqual(result["mobile"], [])
        self.assertEqual(result["national_id"], [])
        self.assertEqual(result["credit_card"], [])

    def test_detect_pii_none_input(self):
        result = detect_pii_in_text(None)
        self.assertEqual(sum(len(v) for v in result.values()), 0)

    # --- analyze_data_exposure（finding 結構與警示）---
    def _page_input(self, html: str) -> PageAnalysisInput:
        return PageAnalysisInput(
            url="https://example.com/",
            final_url="https://example.com/",
            title="測試",
            html=html,
            headers={},
            element_boxes={},
        )

    def test_analyze_data_exposure_returns_no_finding_on_clean_page(self):
        findings = analyze_data_exposure(self._page_input("<p>沒有任何個資的頁面</p>"))
        self.assertEqual(findings, [])

    def test_analyze_data_exposure_finding_contains_warning_prefix(self):
        # 警示文字必須出現在 description 開頭，提醒報告閱讀者責任
        findings = analyze_data_exposure(self._page_input("<p>a@b.com</p>"))
        self.assertEqual(len(findings), 1)
        self.assertTrue(findings[0]["description"].startswith("⚠️ 此項目顯示原始個資"))

    def test_analyze_data_exposure_finding_includes_raw_pii_in_evidence(self):
        # 依使用者要求，evidence 顯示原始個資（不遮罩）
        findings = analyze_data_exposure(self._page_input("<p>a@b.com 0912345678</p>"))
        self.assertEqual(len(findings), 1)
        self.assertIn("a@b.com", findings[0]["evidence"])
        self.assertIn("0912345678", findings[0]["evidence"])

    # --- 分級（2026-09-28 報告審查：看到 Email 不等於外洩）---
    def test_verified_national_id_is_high_severity(self):
        findings = analyze_data_exposure(self._page_input("<p>身分證 A123456789</p>"))
        self.assertEqual(findings[0]["category"], "security")
        self.assertEqual(findings[0]["severity"], "high")
        self.assertEqual(findings[0]["rule_id"], "SECURITY_PII_8B24BB8B28")
        self.assertIn("assessment", findings[0]["evidence_json"])

    def test_external_domain_email_is_medium_personal_contact(self):
        findings = analyze_data_exposure(self._page_input("<p>a@gmail.com</p>"))
        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0]["severity"], "medium")
        self.assertEqual(findings[0]["rule_id"], "security-pii-personal-contact")

    def test_placeholder_sample_number_is_not_pii(self):
        # 註冊表單的填寫範例不是任何人的資料（2026-10-06 ntubimdbirc.tw 註冊頁）
        page = self._page_input('<input type="text" placeholder="e.g.0911-222-333">')
        self.assertEqual(analyze_data_exposure(page), [])

    def test_organizational_email_is_only_info(self):
        # 信箱名稱就是網站名稱（上層機構網域）或角色信箱，屬對外窗口而非個人資料
        page = self._page_input(
            "<p>聯絡 example@school.edu.tw 或 service@partner.org</p>"
        )
        findings = analyze_data_exposure(page)
        self.assertEqual([f["severity"] for f in findings], ["info"])

    def test_site_domain_or_mailto_email_is_only_info(self):
        page = self._page_input(
            '<p>承辦人 staff@example.com</p><a href="mailto:help@other.org">help@other.org</a>'
        )
        findings = analyze_data_exposure(page)
        self.assertEqual([f["severity"] for f in findings], ["info"])
        self.assertEqual(findings[0]["rule_id"], "security-pii-public-contact")

    def test_email_hidden_in_html_comment_is_not_treated_as_public(self):
        findings = analyze_data_exposure(
            self._page_input("<p>hi</p><!-- owner: staff@example.com -->")
        )
        self.assertEqual([f["severity"] for f in findings], ["medium"])

    def test_registrable_domain_handles_second_level_tlds(self):
        from apps.scans.scanners import _registrable_domain

        self.assertEqual(_registrable_domain("imd.ntub.edu.tw"), "ntub.edu.tw")
        self.assertEqual(_registrable_domain("www.example.com"), "example.com")

    def test_analyze_page_includes_data_exposure_when_pii_present(self):
        # PII 偵測必須被掛進 analyze_page pipeline
        page_input = self._page_input(
            "<html><body><h1>x</h1><p>contact: alice@example.tw 0912345678</p></body></html>"
        )
        findings = analyze_page(page_input)
        titles = {f["title"] for f in findings}
        self.assertIn("頁面出現個人聯絡資料 (PII)", titles)


class GeoFastScannerTests(APITestCase):
    def _page_input(self, html: str, html_only: str = "") -> PageAnalysisInput:
        return PageAnalysisInput(
            url="https://example.com/",
            final_url="https://example.com/",
            title="測試頁面",
            html=html,
            headers={},
            element_boxes={},
            html_only=html_only,
        )

    def test_geo_fast_flags_js_dependent_content(self):
        rendered = "<html><body><main>" + "內容文字資料" * 200 + "</main></body></html>"
        raw = "<html><body><div id='root'></div></body></html>"
        page_input = self._page_input(rendered, html_only=raw)
        findings = analyze_geo_fast(page_input, parse_html_signals(rendered))

        self.assertIn("accessible", {finding["impact_area"] for finding in findings})

    def test_geo_fast_skips_accessible_check_for_ssr_page(self):
        html = "<html><body><main>" + "內容文字資料" * 200 + "</main></body></html>"
        page_input = self._page_input(html, html_only=html)
        findings = analyze_geo_fast(page_input, parse_html_signals(html))

        self.assertNotIn("accessible", {finding["impact_area"] for finding in findings})

    def test_geo_fast_flags_missing_semantic_landmark(self):
        html = "<html><body><div>沒有 main 標籤</div></body></html>"
        page_input = self._page_input(html)
        findings = analyze_geo_fast(page_input, parse_html_signals(html))

        self.assertIn("structured", {finding["impact_area"] for finding in findings})

    def test_geo_fast_accepts_semantic_landmark(self):
        html = "<html><body><main>主要內容</main></body></html>"
        page_input = self._page_input(html)
        findings = analyze_geo_fast(page_input, parse_html_signals(html))

        self.assertNotIn("structured", {finding["impact_area"] for finding in findings})

    def test_geo_fast_flags_long_paragraph(self):
        html = f"<html><body><main><p>{'字' * 1500}</p></main></body></html>"
        page_input = self._page_input(html)
        findings = analyze_geo_fast(page_input, parse_html_signals(html))

        self.assertIn("trim", {finding["impact_area"] for finding in findings})

    def test_analyze_site_signals_flags_missing_llms_txt(self):
        findings = analyze_site_signals({"llms_txt_found": False, "blocked_ai_crawlers": []})

        self.assertIn("網站未提供 llms.txt", {finding["title"] for finding in findings})

    def test_analyze_site_signals_flags_blocked_ai_search_bots(self):
        from apps.scans.ai_bots import analyze_policy

        policy = analyze_policy("User-agent: OAI-SearchBot\nDisallow: /\n")
        findings = analyze_site_signals({"llms_txt_found": True, "ai_bot_policy": policy})

        self.assertEqual(
            [f["rule_id"] for f in findings], ["geo-ai-search-bots-blocked"]
        )

    def test_blocking_only_training_bots_is_not_a_problem(self):
        from apps.scans.ai_bots import analyze_policy

        policy = analyze_policy("User-agent: GPTBot\nDisallow: /\n")
        findings = analyze_site_signals({"llms_txt_found": True, "ai_bot_policy": policy})

        self.assertEqual(findings, [])

    def test_analyze_site_signals_clean_site_has_no_findings(self):
        findings = analyze_site_signals({"llms_txt_found": True, "blocked_ai_crawlers": []})

        self.assertEqual(findings, [])


class RerunScanCommandTests(APITestCase):
    def test_rerun_scan_regenerates_findings_from_saved_pages(self):
        user = get_user_model().objects.create_user(
            username="replayer",
            email="replay@example.com",
            password="safe-test-password",
        )
        scan_job = ScanJob.objects.create(
            user=user,
            original_url="https://example.com/",
            normalized_url="https://example.com/",
            origin="https://example.com",
            status=ScanJob.Status.COMPLETED,
        )
        Page.objects.create(
            scan_job=scan_job,
            url="https://example.com/",
            final_url="https://example.com/",
            origin="https://example.com",
            status_code=200,
            title="範例頁",
            html="<html><body><h1>主標題</h1></body></html>",
            html_only_text="<html><body><h1>主標題</h1></body></html>",
            headers={},
            element_boxes={},
        )

        out = StringIO()
        call_command("rerun_scan", scan_job.id, stdout=out)

        self.assertIn("重新掃描", out.getvalue())
        self.assertGreater(Finding.objects.filter(scan_job=scan_job).count(), 0)

    def test_rerun_scan_raises_on_missing_scan_job(self):
        with self.assertRaises(CommandError):
            call_command("rerun_scan", 999999)


class PageTypeRoutingTests(APITestCase):
    """掃描器對 admin 後台與二進位資源的路由判斷：跳過索引面向、保留安全檢查。"""

    def test_is_binary_resource_recognizes_common_downloadables(self):
        self.assertTrue(is_binary_resource("https://example.com/media/app.apk"))
        self.assertTrue(is_binary_resource("https://example.com/files/manual.pdf"))
        self.assertTrue(is_binary_resource("https://example.com/release/build.zip"))
        self.assertTrue(is_binary_resource("https://example.com/assets/logo.png"))

    def test_is_binary_resource_treats_html_pages_as_non_binary(self):
        self.assertFalse(is_binary_resource("https://example.com/"))
        self.assertFalse(is_binary_resource("https://example.com/product/1"))
        self.assertFalse(is_binary_resource("https://example.com/team"))

    def test_is_admin_path_recognizes_common_backend_routes(self):
        self.assertTrue(is_admin_path("https://example.com/admin"))
        self.assertTrue(is_admin_path("https://example.com/admin/login"))
        self.assertTrue(is_admin_path("https://example.com/wp-admin/users.php"))
        self.assertTrue(is_admin_path("https://example.com/dashboard/reports"))

    def test_is_admin_path_skips_frontend_routes(self):
        self.assertFalse(is_admin_path("https://example.com/"))
        self.assertFalse(is_admin_path("https://example.com/administrator-info"))  # 非 admin 前綴
        self.assertFalse(is_admin_path("https://example.com/team"))

    def test_analyze_page_skips_seo_findings_for_admin_login(self):
        # admin/login 缺 H1、缺 description、缺 JSON-LD 都不該被列為 SEO/GEO/AEO 問題
        page_input = PageAnalysisInput(
            url="https://example.com/admin/login",
            final_url="https://example.com/admin/login",
            title="Login",
            html="<html><body><form method='post'><input name='user'></form></body></html>",
            headers={},
            element_boxes={},
        )

        findings = analyze_page(page_input)
        categories = {finding["category"] for finding in findings}

        self.assertNotIn("seo", categories)
        self.assertNotIn("aeo", categories)
        self.assertNotIn("geo", categories)

    def test_analyze_page_keeps_security_findings_for_admin_login(self):
        # 後台登入的 CSRF/安全頭部反而更需要被檢查，不可被跳過
        page_input = PageAnalysisInput(
            url="https://example.com/admin/login",
            final_url="https://example.com/admin/login",
            title="Login",
            html="<html><body><form method='post'><input name='user'></form></body></html>",
            headers={},
            element_boxes={},
        )

        findings = analyze_page(page_input)
        security_titles = {f["title"] for f in findings if f["category"] == "security"}

        self.assertIn("表單可能缺少 CSRF token", security_titles)

    def test_analyze_page_only_runs_security_for_binary_resource(self):
        # APK 連結沒有 HTML 內容，不該被加上 H1/JSON-LD/FAQPage 等建議。
        # HTTPS/HSTS/CSP 等站台層級標頭檢查已搬到 analyze_security_site_level()
        # （對整批頁面只評估一次），這裡改用一個帶表單的頁面驗證 analyze_page()
        # 仍會執行 per-page 的 CSRF 檢查（屬各頁面自己的問題，不能去重站台層級）。
        page_input = PageAnalysisInput(
            url="https://example.com/downloads/app.apk",
            final_url="https://example.com/downloads/app.apk",
            title="",
            html="<form method='post'><input name='q'></form>",
            headers={},
            element_boxes={},
        )

        findings = analyze_page(page_input)
        categories = {finding["category"] for finding in findings}

        self.assertNotIn("seo", categories)
        self.assertNotIn("aeo", categories)
        self.assertNotIn("geo", categories)
        self.assertIn("security", categories)


class CsrfStateChangingFormTests(APITestCase):
    """2026-10-10：只查會改變狀態的表單。

    GOV.UK 實測站內搜尋（GET）被判缺 CSRF token，出現在 8 頁。
    """

    def _security_titles(self, html):
        page_input = PageAnalysisInput(
            url="https://example.com/", final_url="https://example.com/", title="t",
            html=html, headers={}, element_boxes={},
        )
        return {f["title"] for f in analyze_page(page_input) if f["category"] == "security"}

    def test_get_search_form_is_not_flagged(self):
        self.assertNotIn("表單可能缺少 CSRF token", self._security_titles(
            "<form action='/search'><input name='q'><button>搜尋</button></form>"
        ))
        self.assertNotIn("表單可能缺少 CSRF token", self._security_titles(
            "<form method='GET' action='/search'><input name='q'></form>"
        ))

    def test_post_or_password_form_without_token_is_flagged(self):
        self.assertIn("表單可能缺少 CSRF token", self._security_titles(
            "<form method='post' action='/contact'><input name='email'></form>"
        ))
        # 沒寫 method、由 JavaScript 送出的登入表單
        self.assertIn("表單可能缺少 CSRF token", self._security_titles(
            "<form><input name='user'><input type='password' name='pw'></form>"
        ))

    def test_post_form_with_token_is_not_flagged(self):
        self.assertNotIn("表單可能缺少 CSRF token", self._security_titles(
            "<form method='post'><input type='hidden' name='csrfmiddlewaretoken' value='x'></form>"
        ))


class AeoFaqHeuristicTests(APITestCase):
    """FAQPage 建議邏輯：必須真有 FAQ 結構訊號才建議補 Schema，避免機械化誤判。"""

    def _page_input(self, html: str) -> PageAnalysisInput:
        return PageAnalysisInput(
            url="https://example.com/",
            final_url="https://example.com/",
            title="範例頁",
            html=html,
            headers={},
            element_boxes={},
        )

    def test_detect_faq_structure_recognizes_dl(self):
        self.assertTrue(detect_faq_structure("<dl><dt>Q</dt><dd>A</dd></dl>", dl_count=1))

    def test_detect_faq_structure_recognizes_details_tag(self):
        self.assertTrue(detect_faq_structure("<details><summary>Q</summary>A</details>", 0))

    def test_detect_faq_structure_recognizes_faq_class(self):
        self.assertTrue(detect_faq_structure("<section class='faq'>Q&A</section>", 0))

    def test_detect_faq_structure_returns_false_for_plain_content(self):
        self.assertFalse(detect_faq_structure("<p>產品介紹</p>", 0))

    def test_analyze_aeo_does_not_flag_pure_question_tone_without_faq_structure(self):
        # 報告中產品描述常用「能幫你做什麼」「如何使用」等問句，但沒有 FAQ 結構，
        # 此時建議補 FAQPage Schema 是機械化誤判，應改提示「先整理結構」。
        html = (
            "<html><body>"
            "<p>AI 智慧眼鏡能幫你做什麼？我們提供視障導航。</p>"
            "<p>如何使用？戴上即可。</p>"
            "<p>為何選擇我們？因為穩定。</p>"
            "<p>怎麼購買？至產品頁。</p>"
            "</body></html>"
        )
        findings = analyze_aeo(self._page_input(html), parse_html_signals(html))
        titles = {finding["title"] for finding in findings}

        self.assertNotIn("問答內容缺少 FAQPage 或 HowTo 結構化資料", titles)

    def test_analyze_aeo_no_longer_demands_faqpage_schema(self):
        # 2026-09-28 起不再因為沒有 FAQPage／HowTo 就要求補標記：Google 已停止顯示 FAQ
        # 複合搜尋結果，補標記不等於答案更容易被引用；AEO 改看內容能否回答問題。
        html = (
            "<html><body>"
            "<dl>"
            "<dt>什麼是視障導航？</dt><dd>用 AI 辨識環境引導視障者。</dd>"
            "<dt>如何配戴？</dt><dd>像一般眼鏡戴上即可。</dd>"
            "<dt>為何選 Argus？</dt><dd>準確度高。</dd>"
            "<dt>怎麼充電？</dt><dd>Type-C 充電。</dd>"
            "</dl>"
            "</body></html>"
        )
        findings = analyze_aeo(self._page_input(html), parse_html_signals(html))

        self.assertFalse([f for f in findings if "FAQPage" in f["title"]])

    def test_analyze_aeo_ignores_low_question_density_text(self):
        # 內文只出現 1-2 個常用字（什麼/如何），不應觸發任何 AEO finding
        html = "<html><body><p>了解產品是什麼，並如何運作。</p></body></html>"
        findings = analyze_aeo(self._page_input(html), parse_html_signals(html))

        self.assertEqual(findings, [])

    def test_analyze_aeo_ignores_question_text_in_html_attributes(self):
        # 問句訊號只計算可見文字；HTML 屬性或 tag 中的字串不該被計入
        html = (
            "<html><body>"
            "<div class='how-to-image' data-tip='如何 什麼 為何 怎麼 ?'>"
            "<p>純粹的內容，不是問答。</p>"
            "</div></body></html>"
        )
        findings = analyze_aeo(self._page_input(html), parse_html_signals(html))

        self.assertEqual(findings, [])


class EstimateScanTests(APITestCase):
    def setUp(self):
        from django.contrib.auth import get_user_model
        User = get_user_model()
        self.user = User.objects.create_user(username="est@test.com", password="pw")
        from rest_framework_simplejwt.tokens import RefreshToken
        self.token = str(RefreshToken.for_user(self.user).access_token)

    def test_estimate_requires_auth(self):
        from django.test import Client
        c = Client()
        resp = c.post(
            "/api/estimate/",
            {"url": "https://example.com"},
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 401)

    def test_estimate_rejects_invalid_url(self):
        from django.test import Client
        c = Client()
        resp = c.post(
            "/api/estimate/",
            {"url": "not-a-url"},
            content_type="application/json",
            HTTP_AUTHORIZATION=f"Bearer {self.token}",
        )
        self.assertEqual(resp.status_code, 400)

    def test_estimate_rejects_localhost(self):
        from django.test import Client
        c = Client()
        resp = c.post(
            "/api/estimate/",
            {"url": "http://localhost/"},
            content_type="application/json",
            HTTP_AUTHORIZATION=f"Bearer {self.token}",
        )
        self.assertEqual(resp.status_code, 400)

    @override_settings(ARGUS_COIN_PER_CATEGORY=7)
    def test_estimate_uses_billing_coin_per_category_setting(self):
        resp = self.client.post(
            "/api/estimate/",
            {"url": "https://example.com/", "max_pages": 2},
            format="json",
            HTTP_AUTHORIZATION=f"Bearer {self.token}",
        )

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.data["estimated_pages"], 2)
        # 未帶 categories＝五維全選：2 頁 × 5 維 × 7 = 70
        self.assertEqual(resp.data["estimated_cost"], 70)
        self.assertEqual(resp.data["confidence"], "maximum")
        self.assertEqual(resp.data["method"], "billing_cap")

    def test_estimate_counts_selected_categories(self):
        resp = self.client.post(
            "/api/estimate/",
            {"url": "https://example.com/", "max_pages": 10, "categories": ["seo", "geo"]},
            format="json",
            HTTP_AUTHORIZATION=f"Bearer {self.token}",
        )

        self.assertEqual(resp.status_code, 200)
        # 10 頁 × 2 維 × 2 coin = 40
        self.assertEqual(resp.data["estimated_cost"], 40)
        self.assertEqual(resp.data["categories"], ["seo", "geo"])

    def test_estimate_rejects_empty_or_unknown_categories(self):
        for bad in ([], ["seo", "bogus"]):
            resp = self.client.post(
                "/api/estimate/",
                {"url": "https://example.com/", "max_pages": 5, "categories": bad},
                format="json",
                HTTP_AUTHORIZATION=f"Bearer {self.token}",
            )
            self.assertEqual(resp.status_code, 400, msg=bad)

    @patch("apps.scans.services.socket.getaddrinfo")
    def test_estimate_does_not_resolve_dns_or_contact_target(self, getaddrinfo):
        resp = self.client.post(
            "/api/estimate/",
            {"url": "https://example.com/", "max_pages": 50},
            format="json",
            HTTP_AUTHORIZATION=f"Bearer {self.token}",
        )

        self.assertEqual(resp.status_code, 200)
        getaddrinfo.assert_not_called()

    def test_estimate_rejects_non_string_url_without_server_error(self):
        resp = self.client.post(
            "/api/estimate/",
            {"url": {}, "max_pages": 50},
            format="json",
            HTTP_AUTHORIZATION=f"Bearer {self.token}",
        )

        self.assertEqual(resp.status_code, 400)
